"""매물 입력 양식(엑셀) 만들기.

'매물입력' 시트에 한 줄에 매물 하나씩 적어 매물 폴더에 저장하면 프로그램이 읽습니다.
'작성방법' 시트는 설명용이라 매물로 읽히지 않습니다.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Optional

from .models import JIMOK_TYPES
from .requirement_parser import ZONING_TYPES

COLUMNS = [
    # (컬럼명, 너비, 필수, 설명, 예시)
    ("매물번호", 12, False, "사무실에서 쓰는 번호. 비워두면 자동으로 만듭니다.", "E-3101"),
    ("소재지", 42, True, "지번 주소 전체 (시/군/읍·면/리 + 지번)", "충청북도 음성군 금왕읍 무극리 210-3"),
    ("지목", 10, False, "전, 답, 임야, 대, 잡종지 … (목록에서 선택)", "대"),
    ("용도지역", 16, False, "계획관리지역, 자연녹지지역 … (목록에서 선택)", "계획관리지역"),
    ("면적(평)", 10, True, "평 단위 숫자. ㎡로 적으려면 컬럼 이름을 '면적(㎡)'으로 바꾸세요.", "250"),
    ("매매가", 14, True, "'2억 3천만원', '2.3억', '23000'(만원) 모두 가능", "2억 3천만원"),
    ("도로", 18, False, "도로 조건. 길이 없으면 '맹지'", "6m 포장도로 접함"),
    ("상태", 10, False, "거래가능 / 보류 / 계약완료 (비우면 거래가능)", "거래가능"),
    ("등록일", 12, False, "2026-10-12 형식", "2026-10-12"),
    ("특이사항", 44, False, "채광·향·경사·주변 시설 등. 검색에 쓰이니 자세히 적을수록 좋아요.",
     "정남향 평지, 채광 우수, 금왕IC 5분"),
    ("연락처", 15, False, "소유자/매도인 연락처 (손님 화면에는 표시되지 않음)", "010-0000-0000"),
]
DATA_ROWS = 300
STATUSES = ["거래가능", "보류", "계약완료"]


def build_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "매물입력"
    head_fill = PatternFill("solid", fgColor="1F5F4A")
    req_fill = PatternFill("solid", fgColor="B3541E")
    for i, (name, width, required, _, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = req_fill if required else head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24

    # 오타를 줄이는 선택 목록 (목록 밖 값도 입력은 가능하도록 경고만)
    lists = wb.create_sheet("목록")
    for r, v in enumerate(JIMOK_TYPES, start=1):
        lists.cell(row=r, column=1, value=v)
    for r, v in enumerate(ZONING_TYPES, start=1):
        lists.cell(row=r, column=2, value=v)
    for r, v in enumerate(STATUSES, start=1):
        lists.cell(row=r, column=3, value=v)
    lists.sheet_state = "hidden"
    col = {name: ws.cell(row=1, column=i).column_letter for i, (name, *_rest) in enumerate(COLUMNS, start=1)}
    for name, ref in (("지목", f"목록!$A$1:$A${len(JIMOK_TYPES)}"),
                      ("용도지역", f"목록!$B$1:$B${len(ZONING_TYPES)}"),
                      ("상태", f"목록!$C$1:$C${len(STATUSES)}")):
        dv = DataValidation(type="list", formula1=ref, allow_blank=True, showErrorMessage=True,
                            errorStyle="warning", errorTitle="목록에 없는 값",
                            error="목록에 없는 값입니다. 그래도 입력하려면 '예'를 누르세요.")
        dv.add(f"{col[name]}2:{col[name]}{DATA_ROWS + 1}")
        ws.add_data_validation(dv)

    guide = wb.create_sheet("작성방법", 0)
    guide.column_dimensions["A"].width = 14
    guide.column_dimensions["B"].width = 8
    guide.column_dimensions["C"].width = 60
    guide.column_dimensions["D"].width = 40
    guide["A1"] = "토지 매물 입력 양식"
    guide["A1"].font = Font(bold=True, size=14)
    guide["A2"] = "1) '매물입력' 시트에 한 줄에 매물 하나씩 적습니다.  2) 주황색 칸(소재지·면적·매매가)은 꼭 채워주세요."
    guide["A3"] = "3) 프로젝트의 '매물' 폴더에 저장하면 됩니다. (이 양식 파일에 바로 적어 저장해도 되고, 웹 관리자 화면에서 올려도 됩니다)"
    for i, title in enumerate(["항목", "필수", "적는 방법", "예시"], start=1):
        c = guide.cell(row=5, column=i, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    for r, (name, _, required, how, example) in enumerate(COLUMNS, start=6):
        guide.cell(row=r, column=1, value=name)
        guide.cell(row=r, column=2, value="필수" if required else "")
        guide.cell(row=r, column=3, value=how).alignment = Alignment(wrap_text=True)
        guide.cell(row=r, column=4, value=example)
    wb.active = 1  # 열었을 때 '매물입력' 시트가 보이도록

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_template(path: Optional[Path] = None) -> Path:
    path = Path(path or "토지매물_입력양식.xlsx")
    path.write_bytes(build_template())
    return path


TEMPLATE_NAME = "토지매물_입력양식.xlsx"
FOLDER_GUIDE = """이 폴더에 토지 매물 엑셀 파일(.xlsx)을 넣으면 프로그램이 읽습니다.

1. '토지매물_입력양식.xlsx' 를 열어 '매물입력' 시트에 한 줄에 매물 하나씩 적고 저장하세요.
   (파일 이름은 바꿔도 되고, 여러 파일로 나눠도 됩니다. 하위 폴더도 읽습니다)
2. 이미 쓰고 계신 엑셀 장부가 있으면 그 파일을 그대로 넣어도 됩니다.
3. 웹 관리자 화면의 '매물 폴더 다시 읽기' 를 누르거나 프로그램을 다시 켜면 반영됩니다.

※ 이 폴더의 매물 파일은 소유자 연락처가 들어 있어 GitHub 에 올라가지 않도록 설정되어 있습니다.
"""


def prepare_data_folder(folder: Path) -> bool:
    """매물 폴더가 없으면 만들고, 매물 파일이 하나도 없으면 입력 양식과 안내문을 넣어 둠. 만들었으면 True."""
    folder = Path(folder)
    created = not folder.exists()
    folder.mkdir(parents=True, exist_ok=True)
    has_data = any(p.suffix.lower() in {".xlsx", ".xlsm", ".xls", ".csv", ".json"}
                   for p in folder.rglob("*") if p.is_file() and not p.name.startswith("~$"))
    if not has_data:
        (folder / TEMPLATE_NAME).write_bytes(build_template())
    guide = folder / "README.txt"
    if not guide.exists():
        guide.write_text(FOLDER_GUIDE, encoding="utf-8")
    return created
