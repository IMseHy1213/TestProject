"""관리자 화면에서 매물 추가·수정.

- 새 매물 : 매물 폴더의 '웹입력/웹입력매물.xlsx' 에 한 줄씩 추가 -> 다른 엑셀 매물과 똑같이 읽힘
            (엑셀로 열어 직접 고쳐도 되고, 백업할 때 매물 폴더만 복사하면 됨)
- 기존 매물: 원본 엑셀은 건드리지 않고, 고친 항목만 DB(listing_edits)에 저장해 매물을 읽을 때마다 덮어씀
            -> 언제든 '원래대로 되돌리기' 가능. 원본 엑셀을 고치면 고치지 않은 항목은 그대로 반영됨.
"""

from __future__ import annotations

import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from .models import JIMOK_TYPES, PYEONG_M2, LandListing
from .normalizer import normalize_jimok, normalize_zoning, parse_date, parse_price_manwon
from .requirement_parser import ZONING_TYPES

WEB_FOLDER = "웹입력"
WEB_FILE = "웹입력매물.xlsx"
WEB_SHEET = "매물입력"
STATUSES = ["거래가능", "보류", "계약완료"]
HIDDEN = "hidden"  # 목록에서 숨김(삭제) 표시

# 고칠 수 있는 항목 (LandListing 속성 이름)
EDITABLE = ("address", "jimok", "zoning", "area_m2", "price_manwon", "road_access", "status",
            "registered_date", "description", "owner_contact")
LABELS = {
    "address": "소재지", "jimok": "지목", "zoning": "용도지역", "area_m2": "면적", "price_manwon": "매매가",
    "road_access": "도로", "status": "상태", "registered_date": "등록일", "description": "특이사항",
    "owner_contact": "연락처", HIDDEN: "숨김",
}
# 웹입력 엑셀의 열 (프로그램이 읽는 열 이름과 같음)
WEB_COLUMNS = ["매물번호", "소재지", "지목", "용도지역", "면적(㎡)", "매매가", "도로", "상태", "등록일", "특이사항",
               "연락처", "입력일시"]
MAX_TEXT = {"address": 200, "road_access": 100, "description": 2000, "owner_contact": 100}


@dataclass
class ListingForm:
    """화면 입력값을 검사해 LandListing 속성 값으로 바꾼 결과."""

    values: dict[str, Any] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    raw: dict[str, str] = field(default_factory=dict)  # 오류가 있을 때 화면에 다시 채워 줄 입력값


def parse_form(form) -> ListingForm:
    """관리자 화면의 매물 입력 폼 -> 검사된 값."""
    keys = ("address", "jimok", "zoning", "area", "area_unit", "price", "road_access", "status",
            "registered_date", "description", "owner_contact")
    raw = {k: (form.get(k) or "").strip() for k in keys}
    out = ListingForm(raw=raw)
    v, err = out.values, out.errors

    v["address"] = " ".join(raw["address"].split())
    if not v["address"]:
        err.append("소재지를 입력하세요. (예: 충청북도 음성군 금왕읍 무극리 210-3)")
    v["jimok"] = normalize_jimok(raw["jimok"]) if raw["jimok"] else ""
    if v["jimok"] and v["jimok"] not in JIMOK_TYPES:
        err.append(f"지목 '{raw['jimok']}' 을 알 수 없습니다. 목록에서 고르세요.")
    v["zoning"] = normalize_zoning(raw["zoning"]) if raw["zoning"] else ""

    area_text = raw["area"].replace(",", "")
    try:
        area = float(area_text)
        if area <= 0:
            raise ValueError
        v["area_m2"] = round(area * PYEONG_M2 if raw["area_unit"] != "m2" else area, 2)
    except ValueError:
        err.append("면적은 숫자로 입력하세요. (예: 250)")

    price = parse_price_manwon(raw["price"])
    if price is None or price <= 0:
        err.append("매매가를 입력하세요. ('2억 3천만원', '2.3억', '23000'(만원) 모두 가능)")
    else:
        v["price_manwon"] = price

    v["road_access"] = raw["road_access"]
    v["status"] = raw["status"] if raw["status"] in STATUSES else "거래가능"
    if raw["registered_date"]:
        d = parse_date(raw["registered_date"])
        if d is None:
            err.append("등록일은 2026-10-14 형식으로 입력하세요.")
        v["registered_date"] = d
    else:
        v["registered_date"] = None
    v["description"] = raw["description"]
    v["owner_contact"] = raw["owner_contact"]
    for k, n in MAX_TEXT.items():
        if len(v.get(k) or "") > n:
            err.append(f"{LABELS[k]}이(가) 너무 깁니다 ({n}자 이하).")
    return out


def form_defaults(listing: Optional[LandListing] = None) -> dict[str, str]:
    """수정 화면에 미리 채울 값 (면적은 평 단위로 보여줌)."""
    if listing is None:
        return {"area_unit": "pyeong", "status": "거래가능", "registered_date": date.today().isoformat()}
    pyeong = listing.area_m2 / PYEONG_M2
    return {
        "address": listing.address, "jimok": listing.jimok, "zoning": listing.zoning,
        "area": f"{pyeong:.2f}".rstrip("0").rstrip("."), "area_unit": "pyeong",
        "price": str(listing.price_manwon), "road_access": listing.road_access, "status": listing.status,
        "registered_date": listing.registered_date.isoformat() if listing.registered_date else "",
        "description": listing.description, "owner_contact": listing.owner_contact,
    }


# ---------------------------------------------------------------------- 기존 매물 수정 (덮어쓰기)
def _jsonable(value: Any) -> Any:
    return value.isoformat() if isinstance(value, date) else value


def _from_json(key: str, value: Any) -> Any:
    if key == "registered_date" and isinstance(value, str):
        return parse_date(value)
    if key == "area_m2" and value is not None:
        return float(value)
    if key == "price_manwon" and value is not None:
        return int(value)
    return value


def diff_fields(original: LandListing, values: dict[str, Any]) -> dict[str, Any]:
    """원본과 다른 항목만 (DB 에 저장할 값)."""
    out = {}
    for k in EDITABLE:
        if k not in values:
            continue
        new, old = values[k], getattr(original, k)
        if k == "area_m2" and new is not None and old is not None and abs(float(new) - float(old)) < 0.05:
            continue
        if (new or None) != (old or None):
            out[k] = _jsonable(new)
    return out


def apply_fields(listing: LandListing, fields: dict[str, Any]) -> None:
    for k, v in fields.items():
        if k in EDITABLE:
            setattr(listing, k, _from_json(k, v))
    if fields:
        listing.extra["관리자수정"] = sorted(LABELS.get(k, k) for k in fields if k != HIDDEN)


def apply_edits(listings: list[LandListing], edits: dict[str, dict]) -> tuple[
        list[LandListing], list[LandListing], dict[str, LandListing]]:
    """매물을 읽을 때마다 관리자가 고친 내용을 덮어씀. (보이는 매물, 숨긴 매물, 고치기 전 원본) 반환."""
    visible, hidden, originals = [], [], {}
    for l in listings:
        e = edits.get(l.listing_id)
        if e:
            originals[l.listing_id] = copy.deepcopy(l)
            apply_fields(l, e)
        (hidden if e and e.get(HIDDEN) else visible).append(l)
    return visible, hidden, originals


def describe_changes(original: LandListing, fields: dict[str, Any]) -> list[str]:
    """'매매가: 2억 → 1억 8천만원' 같은 변경 내역 (관리자 화면 표시용)."""
    from .models import format_price

    def show(k, v):
        if v in (None, ""):
            return "(없음)"
        if k == "price_manwon":
            return format_price(int(v))
        if k == "area_m2":
            return f"{float(v) / PYEONG_M2:,.1f}평"
        return str(v)

    lines = []
    for k, v in fields.items():
        if k == HIDDEN:
            lines.append("목록에서 숨김")
        elif k in EDITABLE:
            lines.append(f"{LABELS[k]}: {show(k, getattr(original, k))} → {show(k, _from_json(k, v))}")
    return lines


# ---------------------------------------------------------------------- 새 매물 (웹입력 엑셀에 추가)
def web_file(data_dir: Path) -> Path:
    return Path(data_dir) / WEB_FOLDER / WEB_FILE


def new_listing_id(existing: set[str]) -> str:
    base = "W-" + datetime.now().strftime("%y%m%d-%H%M%S")
    candidate, n = base, 1
    while candidate in existing:
        n += 1
        candidate = f"{base}-{n}"
    return candidate


def preview_listing(listing_id: str, values: dict[str, Any], path: Path) -> LandListing:
    """저장하기 전에 검사용으로 만들어 보는 매물."""
    return LandListing(listing_id=listing_id, source="web", source_file=str(path),
                       **{k: values[k] for k in EDITABLE if k in values})


def append_listing(data_dir: Path, listing_id: str, values: dict[str, Any]) -> Path:
    """웹입력 엑셀에 한 줄 추가 (없으면 새로 만듦). 엑셀에서 파일을 열어 두면 PermissionError."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    path = web_file(data_dir)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = load_workbook(path)
        ws = wb[WEB_SHEET] if WEB_SHEET in wb.sheetnames else wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = WEB_SHEET
        for i, name in enumerate(WEB_COLUMNS, start=1):
            c = ws.cell(row=1, column=i, value=name)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F5F4A")
            ws.column_dimensions[c.column_letter].width = 42 if name in {"소재지", "특이사항"} else 14
        ws.freeze_panes = "A2"
    header = [c.value for c in ws[1]]
    row = {
        "매물번호": listing_id, "소재지": values["address"], "지목": values.get("jimok", ""),
        "용도지역": values.get("zoning", ""), "면적(㎡)": values["area_m2"], "매매가": values["price_manwon"],
        "도로": values.get("road_access", ""), "상태": values.get("status", "거래가능"),
        "등록일": values["registered_date"].isoformat() if values.get("registered_date") else "",
        "특이사항": values.get("description", ""), "연락처": values.get("owner_contact", ""),
        "입력일시": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    ws.append([_safe_cell(row.get(h, "")) for h in header])
    tmp = path.with_name("~$tmp_" + path.name)  # '~$' 로 시작하는 파일은 매물로 읽지 않음
    wb.save(tmp)
    tmp.replace(path)
    return path


def _safe_cell(value: Any) -> Any:
    """엑셀 수식으로 해석될 수 있는 글자(=, +, -, @ 로 시작)는 앞에 ' 를 붙여 글자로 저장."""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"} and not value[1:2].isdigit():
        return "'" + value
    return value

