"""관리자 매물 추가·수정 테스트."""

import re
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from real_estate_agent.config import Settings
from real_estate_agent.listing_editor import WEB_FILE, WEB_FOLDER
from real_estate_agent.web.app import create_app

SAMPLE = Path(__file__).resolve().parent.parent / "sample_data" / "토지"

NEW = {"address": "충청북도 음성군 감곡면 오향리 777-1", "jimok": "대", "zoning": "계획관리지역",
       "area": "300", "area_unit": "pyeong", "price": "1억 5천만원", "road_access": "6m 포장도로 접함",
       "status": "거래가능", "registered_date": "2026-10-01", "description": "정남향 평지, 채광 우수",
       "owner_contact": "010-1234-5678"}


@pytest.fixture
def admin(tmp_path, monkeypatch):
    monkeypatch.setenv("REA_ADMIN_PASSWORD", "office-pass-1")
    data = tmp_path / "토지"
    shutil.copytree(SAMPLE, data)
    settings = Settings(data_dir=data, db_path=tmp_path / "w.db", log_dir=tmp_path / "logs", llm_providers=["rule"])
    app = create_app(settings)
    app.config["TESTING"] = True
    c = app.test_client()
    c.post("/admin/login", data={"password": "office-pass-1", "csrf": token(c, "/admin")})
    return c, app.config["AGENT"], data, settings


def token(c, path="/admin/listings"):
    return re.search(r'name="csrf" value="([^"]+)"', c.get(path).get_data(as_text=True)).group(1)


def test_listing_pages_need_login(tmp_path, monkeypatch):
    monkeypatch.setenv("REA_ADMIN_PASSWORD", "office-pass-1")
    settings = Settings(data_dir=SAMPLE, db_path=tmp_path / "w.db", log_dir=tmp_path / "l", llm_providers=["rule"])
    c = create_app(settings).test_client()
    assert c.get("/admin/listings").status_code == 302
    assert c.get("/admin/listings/new").status_code == 302
    t = token(c, "/")
    assert c.post("/admin/listings/new", data={**NEW, "csrf": t}).status_code == 302
    assert c.post("/admin/listings/state", data={"id": "x", "op": "hide", "csrf": t}).status_code == 403


def test_add_new_listing_goes_to_excel_and_search(admin):
    c, agent, data, _ = admin
    before = len(agent.listings())
    r = c.post("/admin/listings/new", data={**NEW, "csrf": token(c)}, follow_redirects=True)
    html = r.get_data(as_text=True)
    assert "새 매물을 추가했습니다" in html and "오향리 777-1" in html
    assert len(agent.listings()) == before + 1
    path = data / WEB_FOLDER / WEB_FILE
    ws = load_workbook(path).active
    assert ws.max_row == 2 and "오향리 777-1" in str(ws.cell(row=2, column=2).value)
    new = next(l for l in agent.listings() if "오향리 777-1" in l.address)
    assert new.price_manwon == 15000 and round(new.area_pyeong) == 300 and new.listing_id.startswith("W-")

    agent.sync()  # 프로그램을 다시 켜도(폴더를 다시 읽어도) 그대로 있음
    assert any("오향리 777-1" in l.address for l in agent.listings())

    # 같은 땅을 또 넣으면 막음
    r = c.post("/admin/listings/new", data={**NEW, "csrf": token(c)})
    assert r.status_code == 400 and "이미 같은 매물" in r.get_data(as_text=True)


def test_new_listing_validation(admin):
    c, agent, _, _ = admin
    bad = {**NEW, "address": "", "area": "abc", "price": ""}
    r = c.post("/admin/listings/new", data={**bad, "csrf": token(c)})
    html = r.get_data(as_text=True)
    assert r.status_code == 400 and "소재지를 입력" in html and "면적은 숫자" in html and "매매가를 입력" in html
    r = c.post("/admin/listings/new", data={**NEW, "price": "1", "csrf": token(c)})  # 평당 0.003만원
    assert r.status_code == 400 and "평당가" in r.get_data(as_text=True)
    r = c.post("/admin/listings/new", data={**NEW, "description": "=HYPERLINK(\"x\")", "csrf": token(c)})
    assert r.status_code == 302  # 수식처럼 보이는 글자는 엑셀에 글자로 저장
    ws = load_workbook(Path(agent.settings.data_dir) / WEB_FOLDER / WEB_FILE).active
    assert str(ws.cell(row=2, column=10).value).startswith("'=")


def test_edit_existing_listing_overrides_without_touching_excel(admin):
    c, agent, data, _ = admin
    target = sorted(agent.listings(), key=lambda l: l.listing_id)[0]
    lid, old_price = target.listing_id, target.price_manwon
    excel_before = {p: p.read_bytes() for p in data.rglob("*.xlsx")}

    page = c.get(f"/admin/listings/edit?id={lid}").get_data(as_text=True)
    assert target.address in page
    form = dict(re.findall(r'name="(\w+)" value="([^"]*)"', page))
    form.update(price=str(old_price - 1000), status="보류", description="가격 조정됨, 급매",
                jimok=target.jimok, area_unit="pyeong", csrf=token(c))
    r = c.post("/admin/listings/edit", data=form, follow_redirects=True)
    assert "수정했습니다" in r.get_data(as_text=True)
    l = agent.find_listing(lid)
    assert l.price_manwon == old_price - 1000 and l.status == "보류" and "급매" in l.description
    assert {p: p.read_bytes() for p in data.rglob("*.xlsx")} == excel_before  # 원본 엑셀은 그대로

    agent.sync()  # 폴더를 다시 읽어도 고친 내용 유지
    l = agent.find_listing(lid)
    assert l.price_manwon == old_price - 1000 and l.extra.get("관리자수정")
    page = c.get(f"/admin/listings/edit?id={lid}").get_data(as_text=True)
    assert "지금까지 고친 내용" in page and "매매가" in page

    c.post("/admin/listings/state", data={"id": lid, "op": "revert", "csrf": token(c)})
    assert agent.find_listing(lid).price_manwon == old_price
    agent.sync()
    assert agent.find_listing(lid).price_manwon == old_price and not agent.storage.listing_edits()


def test_hide_and_show_listing(admin):
    c, agent, _, _ = admin
    lid = agent.listings()[0].listing_id
    n = len(agent.listings())
    c.post("/admin/listings/state", data={"id": lid, "op": "hide", "csrf": token(c)})
    assert len(agent.listings()) == n - 1 and agent.find_listing(lid) is not None
    html = c.get("/admin/listings?show=hidden").get_data(as_text=True)
    assert lid in html and "숨김" in html
    agent.sync()
    assert all(l.listing_id != lid for l in agent.listings())  # 다시 읽어도 숨김 유지
    c.post("/admin/listings/state", data={"id": lid, "op": "show", "csrf": token(c)})
    assert any(l.listing_id == lid for l in agent.listings()) and len(agent.listings()) == n


def test_list_search_and_edit_link_on_result(admin):
    c, agent, _, _ = admin
    some = agent.listings()[0]
    word = some.address.split()[-2]
    html = c.get(f"/admin/listings?q={word}").get_data(as_text=True)
    assert some.listing_id in html and "새 매물 추가" in html
    r = c.post("/search", data={"q": "음성군 2억원대 땅", "csrf": token(c, "/")}, follow_redirects=True)
    assert "이 매물 수정" in r.get_data(as_text=True)
