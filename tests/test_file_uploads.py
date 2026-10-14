"""관리자 매물 파일·폴더 올리기 테스트."""

import io
import re
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook
from werkzeug.datastructures import FileStorage

from real_estate_agent.config import Settings
from real_estate_agent.file_uploads import BACKUP_FOLDER, clean_relative_path, resolve_destination, save_uploads
from real_estate_agent.web.app import create_app

SAMPLE = Path(__file__).resolve().parent.parent / "sample_data" / "토지"


def info_sheet(rows) -> bytes:
    """'항목 | 값' 형식의 매물 정보지 엑셀."""
    wb = Workbook()
    ws = wb.active
    for k, v in rows:
        ws.append([k, v])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ledger(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["매물번호", "소재지", "지목", "면적(평)", "매매가", "특이사항"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def admin(tmp_path, monkeypatch):
    monkeypatch.setenv("REA_ADMIN_PASSWORD", "office-pass-1")
    data = tmp_path / "매물"
    shutil.copytree(SAMPLE, data)
    (data / "음성군").mkdir()
    settings = Settings(data_dir=data, db_path=tmp_path / "w.db", log_dir=tmp_path / "logs", llm_providers=["rule"])
    app = create_app(settings)
    app.config["TESTING"] = True
    c = app.test_client()
    c.post("/admin/login", data={"password": "office-pass-1", "csrf": token(c, "/admin")})
    return c, app.config["AGENT"], data


def token(c, path="/admin/listings"):
    return re.search(r'name="csrf" value="([^"]+)"', c.get(path).get_data(as_text=True)).group(1)


def upload(c, files, **form):
    data = {"csrf": token(c), "back": "listings", **form}
    data["folder"] = [(io.BytesIO(content), name) for name, content in files]
    return c.post("/admin/upload", data=data, content_type="multipart/form-data", follow_redirects=True)


FOLDER = "음성군 감곡면 오향리 501-2 외 3필지(물류부지)"
INFO = [("소재지", "충청북도 음성군 감곡면 오향리 501-2"), ("지목", "공장용지"), ("면적", "3,000평"),
        ("매매가", "15억"), ("특이사항", "물류창고 부지, 대형차 진입 가능")]


def test_folder_upload_keeps_structure_and_becomes_listing(admin):
    c, agent, data = admin
    r = upload(c, [(f"{FOLDER}/{FOLDER}.xlsx", info_sheet(INFO)), (f"{FOLDER}/현장사진.jpg", b"\xff\xd8"),
                   (f"{FOLDER}/~$임시.xlsx", b"x")], dest="음성군")
    html = r.get_data(as_text=True)
    assert "파일 1개 저장" in html and "1개 건너뜀" in html and "새 매물 1건" in html
    assert (data / "음성군" / FOLDER / f"{FOLDER}.xlsx").exists()
    assert not (data / "음성군" / FOLDER / "현장사진.jpg").exists()  # 사진은 저장 안 함
    l = next(l for l in agent.listings() if "오향리 501-2" in l.address)
    assert l.price_manwon == 150000 and "물류" in l.description


def test_multiple_ledger_files_and_new_folder(admin):
    c, agent, data = admin
    before = len(agent.listings())
    r = upload(c, [("진천_매물.xlsx", ledger([["J-1", "충청북도 진천군 진천읍 장관리 11", "전", 300, "1억 2천"],
                                             ["J-2", "충청북도 진천군 문백면 구곡리 22", "답", 500, "9천만원"]])),
                   ("메모.csv", "소재지,면적(평),매매가\n충청북도 진천군 이월면 송림리 33,200,8000\n".encode("utf-8-sig"))],
               dest="__new__", new_dest="진천군")
    assert "파일 2개 저장" in r.get_data(as_text=True)
    assert len(agent.listings()) == before + 3 and (data / "진천군" / "진천_매물.xlsx").exists()
    assert "진천군" in c.get("/admin").get_data(as_text=True)  # 다음부터 '저장할 폴더' 목록에 나옴


def test_same_name_keeps_both_or_replaces_with_backup(admin):
    c, agent, data = admin
    first = ledger([["K-1", "충청북도 음성군 삼성면 덕정리 5", "대", 200, "2억"]])
    upload(c, [("장부.xlsx", first)], dest="음성군")
    r = upload(c, [("장부.xlsx", first)], dest="음성군")  # 기본: 둘 다 보관
    assert "번호를 붙임" in r.get_data(as_text=True) and (data / "음성군" / "장부_2.xlsx").exists()
    (data / "음성군" / "장부_2.xlsx").unlink()

    second = ledger([["K-1", "충청북도 음성군 삼성면 덕정리 5", "대", 200, "1억 8천"]])
    r = upload(c, [("장부.xlsx", second)], dest="음성군", replace="1")
    assert "_백업" in r.get_data(as_text=True)
    backups = list((data / BACKUP_FOLDER).rglob("장부.xlsx"))
    assert len(backups) == 1  # 기존 파일은 _백업 에 보관
    assert agent.find_listing("K-1").price_manwon == 18000  # 새 파일 내용
    assert len([l for l in agent.listings() if l.listing_id.startswith("K-1")]) == 1  # 백업본은 매물로 안 읽음


def test_upload_into_listing_folder_from_edit_page(admin):
    c, agent, data = admin
    upload(c, [(f"{FOLDER}/{FOLDER}.xlsx", info_sheet(INFO))], dest="음성군")
    lid = next(l for l in agent.listings() if "오향리 501-2" in l.address).listing_id
    page = c.get(f"/admin/listings/edit?id={lid}").get_data(as_text=True)
    assert "이 매물의 파일" in page and f"음성군/{FOLDER}" in page

    new_info = [(k, "13억" if k == "매매가" else v) for k, v in INFO]
    r = c.post("/admin/upload", data={"csrf": token(c), "back": "edit", "id": lid, "replace": "1",
                                      "dest": "음성군",  # 무시되고 이 매물의 폴더에 저장
                                      "files": [(io.BytesIO(info_sheet(new_info)), f"{FOLDER}.xlsx")]},
               content_type="multipart/form-data", follow_redirects=True)
    assert "이 매물의 파일" in r.get_data(as_text=True)  # 수정 화면으로 돌아옴
    assert agent.find_listing(lid).price_manwon == 130000


def test_paths_cannot_escape_data_folder(admin, tmp_path):
    c, agent, data = admin
    assert clean_relative_path("../../evil.xlsx") is None
    assert clean_relative_path("a/../../b.xlsx") is None
    assert str(clean_relative_path('음성군\\폴더:이름?/파일.xlsx')) == "음성군/폴더이름/파일.xlsx"
    with pytest.raises(ValueError):
        resolve_destination(data, "../밖")
    with pytest.raises(ValueError):
        resolve_destination(data, BACKUP_FOLDER)
    r = upload(c, [("../../evil.xlsx", ledger([]))], dest="음성군")
    assert "올리지 못했습니다" in r.get_data(as_text=True)
    assert not (tmp_path / "evil.xlsx").exists() and not (data.parent / "evil.xlsx").exists()
    with pytest.raises(ValueError):
        save_uploads(data, [FileStorage(io.BytesIO(b"x"), "a.xlsx")], dest="../../")


def test_upload_requires_admin(tmp_path, monkeypatch):
    monkeypatch.setenv("REA_ADMIN_PASSWORD", "office-pass-1")
    settings = Settings(data_dir=SAMPLE, db_path=tmp_path / "w.db", log_dir=tmp_path / "l", llm_providers=["rule"])
    c = create_app(settings).test_client()
    r = c.post("/admin/upload", data={"csrf": token(c, "/"), "files": [(io.BytesIO(b"x"), "a.xlsx")]},
               content_type="multipart/form-data")
    assert r.status_code == 403


def test_folder_paths_field_restores_structure(admin):
    """브라우저가 파일 이름만 보내도 화면이 보낸 경로로 폴더 구조를 살림."""
    c, agent, data = admin
    r = c.post("/admin/upload", data={"csrf": token(c), "back": "listings", "dest": "음성군",
                                      "folder": [(io.BytesIO(info_sheet(INFO)), f"{FOLDER}.xlsx")],
                                      "folder_paths": [f"{FOLDER}/{FOLDER}.xlsx"]},
               content_type="multipart/form-data", follow_redirects=True)
    assert "새 매물 1건" in r.get_data(as_text=True)
    assert (data / "음성군" / FOLDER / f"{FOLDER}.xlsx").exists()


def test_office_number_from_excel_is_shown(admin):
    c, agent, data = admin
    with_no = [("매물번호", "음성-0715")] + INFO
    upload(c, [(f"{FOLDER}/{FOLDER}.xlsx", info_sheet(with_no))], dest="음성군")
    l = next(l for l in agent.listings() if "오향리 501-2" in l.address)
    assert l.office_no == "음성-0715"

    # 필지 목록의 '번호'(1, 2, 3 순번)는 매물번호로 쓰지 않음
    other = "음성군 생극면 방축리 88"
    wb = Workbook()
    ws = wb.active
    ws.append(["번호", "소재지", "지목", "면적(㎡)", "금액"])
    ws.append([1, "충청북도 음성군 생극면 방축리 88", "전", 1000, 5000])
    ws.append([2, "충청북도 음성군 생극면 방축리 89", "전", 1200, 6000])
    buf = io.BytesIO()
    wb.save(buf)
    upload(c, [(f"{other}/{other}.xlsx", buf.getvalue())], dest="음성군")
    l2 = next(l for l in agent.listings() if "방축리 88" in l.address)
    assert l2.office_no == "" and l2.listing_id.startswith("F-")

    ledger_listing = agent.find_listing("T-1004")  # 장부 파일의 매물번호 열
    assert ledger_listing.office_no == "T-1004"

    html = c.get(f"/admin/listings?q=음성-0715").get_data(as_text=True)
    assert "음성-0715" in html and l.listing_id in html  # 매물번호로 찾기 + 프로그램 번호도 표시
